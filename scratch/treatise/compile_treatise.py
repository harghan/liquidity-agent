import os
import markdown
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

VOLUMES = [
    "/Users/harshaghandikota/Liquidity Agent/docs/treatise/vol1.md",
    "/Users/harshaghandikota/Liquidity Agent/docs/treatise/vol2.md",
    "/Users/harshaghandikota/Liquidity Agent/docs/treatise/vol3.md",
    "/Users/harshaghandikota/Liquidity Agent/docs/treatise/vol4.md",
    "/Users/harshaghandikota/Liquidity Agent/docs/treatise/vol5.md",
]

def assemble_master_markdown(out_path):
    full_text = []
    full_text.append("# EPISTEMIC SOVEREIGNTY IN THE AGE OF EVENT DERIVATIVES\n")
    full_text.append("## A Master Treatise on Macro Civilization, Market Microstructure, and Autonomous Agentic Capital\n\n")
    full_text.append("**By Harsha Ghandikota**  \n*PRISM Technologies Inc. — Quantitative Systems & Civilization Architecture*  \n*October 2026*\n\n---\n\n")
    
    for v_path in VOLUMES:
        if os.path.exists(v_path):
            with open(v_path, "r", encoding="utf-8") as f:
                full_text.append(f.read())
                full_text.append("\n\n<div style='page-break-after: always;'></div>\n\n---\n\n")
    
    combined = "\n".join(full_text)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(combined)
    print(f"Master markdown created at: {out_path}")
    return combined

def assemble_master_html(md_path, html_path):
    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()
    
    # Custom HTML template with Newsreader / Garamond publication styling
    html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Epistemic Sovereignty in the Age of Event Derivatives — Harsha Ghandikota</title>
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/katex@0.16.8/dist/katex.min.css">
<style>
  @page {{
    size: letter;
    margin: 1.0in 0.9in 1.0in 0.9in;
    @bottom-center {{
      content: counter(page);
      font-family: 'Newsreader', 'EB Garamond', Georgia, serif;
      font-size: 9pt;
      color: #64748b;
    }}
  }}
  body {{
    font-family: 'Newsreader', 'EB Garamond', Georgia, serif;
    font-size: 11pt;
    line-height: 1.62;
    color: #1e293b;
    background-color: #ffffff;
    max-width: 820px;
    margin: 0 auto;
    padding: 20px;
  }}
  h1 {{
    font-family: 'Inter', -apple-system, sans-serif;
    font-size: 24pt;
    font-weight: 800;
    line-height: 1.15;
    color: #0f172a;
    letter-spacing: -0.03em;
    margin-top: 40px;
    margin-bottom: 12px;
    border-bottom: 2px solid #0f172a;
    padding-bottom: 8px;
    page-break-before: always;
  }}
  h1:first-of-type {{
    page-break-before: avoid;
    margin-top: 20px;
  }}
  h2 {{
    font-family: 'Inter', -apple-system, sans-serif;
    font-size: 15pt;
    font-weight: 700;
    color: #1e293b;
    margin-top: 28px;
    margin-bottom: 10px;
    letter-spacing: -0.02em;
  }}
  h3 {{
    font-family: 'Inter', -apple-system, sans-serif;
    font-size: 12pt;
    font-weight: 600;
    color: #334155;
    margin-top: 20px;
    margin-bottom: 8px;
  }}
  p {{
    margin-bottom: 12px;
    text-align: justify;
    text-justify: inter-word;
  }}
  blockquote {{
    margin: 20px 0;
    padding: 14px 20px;
    background-color: #f8fafc;
    border-left: 3px solid #0f172a;
    font-style: italic;
    color: #334155;
  }}
  pre, code {{
    font-family: 'JetBrains Mono', monospace;
    font-size: 8.5pt;
    background-color: #f8fafc;
    color: #0f172a;
  }}
  pre {{
    padding: 12px;
    border: 1px solid #e2e8f0;
    border-radius: 4px;
    overflow-x: auto;
    line-height: 1.35;
    margin: 16px 0;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    margin: 20px 0;
    font-size: 9.5pt;
  }}
  th {{
    border-top: 1.5px solid #0f172a;
    border-bottom: 1.5px solid #0f172a;
    padding: 8px 10px;
    text-align: left;
    font-family: 'Inter', sans-serif;
    font-weight: 600;
    color: #0f172a;
    background-color: #f8fafc;
  }}
  td {{
    border-bottom: 1px solid #e2e8f0;
    padding: 8px 10px;
    color: #1e293b;
  }}
  tr:last-child td {{
    border-bottom: 1.5px solid #0f172a;
  }}
  hr {{
    border: 0;
    border-top: 1px solid #e2e8f0;
    margin: 30px 0;
  }}
  .meta-block {{
    font-family: 'Inter', sans-serif;
    font-size: 10pt;
    color: #64748b;
    margin-bottom: 30px;
  }}
  .page-break {{
    page-break-after: always;
  }}
</style>
</head>
<body>
{markdown.markdown(md_content, extensions=['tables', 'fenced_code'])}
</body>
</html>
"""
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_template)
    print(f"Master HTML written to: {html_path}")

def generate_master_docx(md_path, docx_path):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)

    FONT = "Georgia"
    
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        
        if line_s.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(line_s[2:])
            r.font.name = FONT
            r.font.size = Pt(20)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
        elif line_s.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(line_s[3:])
            r.font.name = FONT
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 41, 59)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif line_s.startswith("### "):
            p = doc.add_paragraph()
            r = p.add_run(line_s[4:])
            r.font.name = FONT
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(51, 65, 85)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif line_s.startswith("> "):
            p = doc.add_paragraph()
            r = p.add_run(line_s[2:])
            r.font.name = FONT
            r.font.size = Pt(10)
            r.font.italic = True
            r.font.color.rgb = RGBColor(71, 85, 105)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(6)
        elif line_s.startswith("---"):
            p = doc.add_paragraph()
            r = p.add_run("____________________________________________________")
            r.font.color.rgb = RGBColor(226, 232, 240)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
        elif line_s.startswith("```"):
            continue
        else:
            p = doc.add_paragraph()
            r = p.add_run(line_s)
            r.font.name = FONT
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(30, 41, 59)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.25

    doc.save(docx_path)
    print(f"Master DOCX saved to: {docx_path}")

if __name__ == "__main__":
    md_out = "/Users/harshaghandikota/Liquidity Agent/docs/treatise/PRISM_Macro_Civilization_Master_Treatise.md"
    html_out = "/Users/harshaghandikota/Liquidity Agent/docs/treatise/PRISM_Macro_Civilization_Master_Treatise.html"
    docx_out = "/Users/harshaghandikota/Liquidity Agent/docs/treatise/PRISM_Macro_Civilization_Master_Treatise.docx"
    
    assemble_master_markdown(md_out)
    assemble_master_html(md_out, html_out)
    generate_master_docx(md_out, docx_out)
